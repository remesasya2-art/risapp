"""Que las fotos de los comprobantes no viajen cuando nadie las pidió.

DONDE ESTAN LAS FOTOS

    El comprobante de una recarga y los del pago de un retiro se guardan
    INLINE, como data URL en base64, adentro del documento de `transactions`.
    Una foto de celular de 500 KB pesa unos 667 KB así, y un retiro completado
    lleva una LISTA de fotos, no una.

QUE PASABA, MEDIDO CORRIENDOLO

    Veintiuna consultas a `transactions` no llevaban proyección. La peor era
    la de exportar:

        db.transactions.find({}, {"_id": 0, "proof_image": 0}).to_list(10000)

    Con cincuenta filas de una foto en singular y dos en plural: 64 MB en
    memoria para escribir NUEVE columnas de un Excel. Al tope de la ruta —diez
    mil filas— unos 12 GB. El proceso de Railway muere mucho antes y se lleva
    puesta la app para todos mientras se reinicia.

Y POR QUE ESA PROYECCION NO ALCANZABA

    Era una lista de lo PROHIBIDO y nombraba `proof_image`, el campo viejo.
    Después se agregó `proof_images` —en plural, una lista— y nadie se acordó
    de agregarlo ahí. Es exactamente la forma de fallar contra la que avisa la
    regla del proyecto.

LO QUE ESTE ARCHIVO VIGILA

    1. Que toda consulta de LISTA a `transactions` lleve proyección.
    2. Que si la proyección es de lo prohibido, excluya TODOS los campos de
       fotos, no el que alguien se acordó.
    3. Que la ruta de exportar no se traiga ninguna foto, comprobado con
       fotos de verdad en la base y midiendo lo que vuelve.

POR QUE `find_one` QUEDA AFUERA

    Buscar UNA transacción por su id es, casi siempre, la pantalla de detalle:
    ahí la foto es justo lo que se viene a ver. Traer una foto cuando se pidió
    una transacción no es el problema; traer mil cuando se pidió una lista, sí.
"""
import ast
import asyncio
import os
import pathlib
import sys

import pytest

_BACKEND = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
sys.path.insert(0, _BACKEND)

os.environ.setdefault("MONGO_URL", "mongodb://localhost:27017")
os.environ.setdefault("DB_NAME", "ris_test")

mongomock_motor = pytest.importorskip(
    "mongomock_motor",
    reason="mongomock-motor no está instalado: es de test y no va en producción")

from conftest import usar_base                                      # noqa: E402
from services import las_fotos                                      # noqa: E402


def corre(coro):
    return asyncio.run(coro)


@pytest.fixture
def base():
    b = mongomock_motor.AsyncMongoMockClient()["ris_las_fotos"]
    usar_base(b)
    yield b


# Una foto de celular de 500 KB, guardada como data URL.
FOTO = "data:image/jpeg;base64," + ("A" * 667_000)


# ══════════════════════════════════════════════════════════════════════════
# 1. Ninguna consulta de lista sin proyección
# ══════════════════════════════════════════════════════════════════════════

def _consultas_de_lista():
    """(archivo, línea, proyección) de cada `transactions.find(...)` del código."""
    raiz = pathlib.Path(_BACKEND)
    for p in sorted(raiz.rglob("*.py")):
        rel = p.relative_to(raiz).as_posix()
        if rel.startswith(("tests/", "venv/", "migrations/", "scripts/")):
            continue
        try:
            arbol = ast.parse(p.read_text())
        except SyntaxError:                                  # pragma: no cover
            continue
        for n in ast.walk(arbol):
            if not (isinstance(n, ast.Call) and isinstance(n.func, ast.Attribute)
                    and n.func.attr == "find"):
                continue
            if not ast.unparse(n.func.value).endswith("transactions"):
                continue
            proyeccion = n.args[1] if len(n.args) >= 2 else None
            yield rel, n.lineno, proyeccion, ast.unparse(n)


def test_ninguna_consulta_de_lista_va_sin_proyeccion():
    sueltas = [f"{rel}:{ln}  {txt[:80]}"
               for rel, ln, proy, txt in _consultas_de_lista() if proy is None]
    assert not sueltas, (
        "Estas consultas a `transactions` traen el documento entero, con las "
        "fotos del comprobante en base64 adentro:\n\n"
        + "\n".join(f"    {x}" for x in sueltas)
        + "\n\nUsá `las_fotos.solo(...)` con los campos que la pantalla "
          "escribe, o `las_fotos.sin_las_fotos()` si el documento se le pasa "
          "entero a la vista."
    )


# `{"_id": 0}` sola NO es una proyección: saca el identificador de Mongo y
# deja pasar todo lo demás, fotos incluidas. Durante un tiempo este archivo la
# contaba como «lista de lo permitido» —no excluye nada fuera de `_id`— y por
# ese agujero pasaban cuatro consultas, entre ellas la exportación a Excel de
# hasta diez mil órdenes: el mismo caso de 12 GB que este archivo describe
# arriba, vuelto a abrir por la puerta de al lado.
#
# Las únicas que pueden traer el documento entero son las que MUESTRAN las
# fotos, y cada una va acá con su motivo. Se identifican por archivo y función,
# no por línea: una línea se corre con cualquier cambio y la excepción
# terminaría tapando otra consulta.
TRAEN_LAS_FOTOS_A_PROPOSITO = {
    ("services/recargas_ves.py", "cola"):
        "la cola de recargas en bolívares del panel muestra el comprobante de "
        "cada fila para aprobarla sin abrir otra pantalla",
    ("services/retiros.py", "cola"):
        "la cola de retiros del panel muestra los comprobantes de pago de "
        "cada fila",
}


def _consultas_con_su_funcion():
    """(archivo, función, línea, proyección) de cada `transactions.find(...)`."""
    raiz = pathlib.Path(_BACKEND)
    for p in sorted(raiz.rglob("*.py")):
        rel = p.relative_to(raiz).as_posix()
        if rel.startswith(("tests/", "venv/", "migrations/", "scripts/")):
            continue
        try:
            arbol = ast.parse(p.read_text())
        except SyntaxError:                                  # pragma: no cover
            continue
        for funcion in ast.walk(arbol):
            if not isinstance(funcion, (ast.FunctionDef, ast.AsyncFunctionDef)):
                continue
            for n in ast.walk(funcion):
                if (isinstance(n, ast.Call) and isinstance(n.func, ast.Attribute)
                        and n.func.attr == "find"
                        and ast.unparse(n.func.value).endswith("transactions")):
                    proyeccion = n.args[1] if len(n.args) >= 2 else None
                    yield rel, funcion.name, n.lineno, proyeccion


def _solo_saca_el_id(proyeccion) -> bool:
    if not isinstance(proyeccion, ast.Dict):
        return False
    claves = [k.value for k in proyeccion.keys if isinstance(k, ast.Constant)]
    return len(claves) == len(proyeccion.keys) and set(claves) <= {"_id"}


def test_una_proyeccion_que_solo_saca_el_id_cuenta_como_ninguna():
    enteras = sorted({f"{rel}:{ln} (en `{fn}`)"
                      for rel, fn, ln, proy in _consultas_con_su_funcion()
                      if _solo_saca_el_id(proy)
                      and (rel, fn) not in TRAEN_LAS_FOTOS_A_PROPOSITO})
    assert not enteras, (
        "Estas consultas usan `{\"_id\": 0}`, que trae la orden ENTERA, con "
        "los comprobantes en base64 adentro:\n\n"
        + "\n".join(f"    {x}" for x in enteras)
        + "\n\nUsá `las_fotos.solo(...)` con los campos que se usan. Si la "
          "pantalla muestra las fotos, agregala a TRAEN_LAS_FOTOS_A_PROPOSITO "
          "con el motivo."
    )


def test_las_excepciones_siguen_existiendo():
    """Una excepción cuya consulta ya no existe tapa la que venga a ocupar su
    lugar. Si se borra o se arregla la consulta, se borra la excepción."""
    vistas = {(rel, fn) for rel, fn, ln, proy in _consultas_con_su_funcion()
              if _solo_saca_el_id(proy)}
    sobran = set(TRAEN_LAS_FOTOS_A_PROPOSITO) - vistas
    assert not sobran, f"excepciones sin consulta: {sorted(sobran)}"


def test_el_recorrido_encuentra_las_consultas_que_dice_vigilar():
    """Sin esto, el de arriba pasaría igual si el recorrido dejara de
    encontrar consultas: cero de cero es verde y no prueba nada."""
    cuantas = sum(1 for _ in _consultas_de_lista())
    assert cuantas >= 15, (
        f"El recorrido sólo vio {cuantas} consultas de lista a "
        f"`transactions`. Se rompió: este archivo no está mirando lo que dice.")


def test_toda_lista_de_lo_prohibido_saca_TODAS_las_fotos():
    """El defecto original: la proyección nombraba una de las tres.

    Agregar un campo de fotos nuevo a `las_fotos.LAS_FOTOS` pone en rojo cada
    lugar que se olvidó de sacarlo. Antes, agregarlo no avisaba en ningún lado.
    """
    incompletas = []
    for rel, ln, proy, txt in _consultas_de_lista():
        if not isinstance(proy, ast.Dict):
            continue                      # `las_fotos.sin_las_fotos()`, o `solo(...)`
        excluidos = {k.value for k, v in zip(proy.keys, proy.values)
                     if isinstance(k, ast.Constant) and isinstance(v, ast.Constant)
                     and v.value == 0}
        if not excluidos - {"_id"}:
            continue                      # es una lista de lo permitido
        faltan = set(las_fotos.LAS_FOTOS) - excluidos
        if faltan:
            incompletas.append(f"{rel}:{ln} le falta sacar {sorted(faltan)}")
    assert not incompletas, (
        "Proyecciones de lo PROHIBIDO que no sacan todas las fotos:\n\n"
        + "\n".join(f"    {x}" for x in incompletas)
        + "\n\nUsá `las_fotos.sin_las_fotos()`, que las nombra a todas en un "
          "solo lugar."
    )


# ══════════════════════════════════════════════════════════════════════════
# 2. Y que se cumpla de verdad, con fotos en la base
# ══════════════════════════════════════════════════════════════════════════

def test_exportar_no_se_trae_ni_una_foto(base):
    """La ruta que podía matar el proceso. Se mide lo que vuelve."""
    async def cuerpo():
        import admin_routes
        admin_routes.db = base
        for i in range(30):
            await base.transactions.insert_one({
                "transaction_id": f"tx{i}", "user_id": "u1", "type": "withdrawal",
                "status": "completed", "proof_image": FOTO,
                "proof_images": [FOTO, FOTO],
            })
        filas = await base.transactions.find({}, las_fotos.solo(
            "transaction_id", "user_id", "type", "status",
            "amount_input", "amount_output", "created_at", "completed_at",
            "beneficiary_data")).to_list(10000)
        peso = sum(len(str(f)) for f in filas)
        assert peso < 30 * 1024, (
            f"treinta filas pesan {peso/1024/1024:.1f} MB: las fotos siguen "
            f"viajando")
        for campo in las_fotos.LAS_FOTOS:
            assert campo not in filas[0], f"«{campo}» llegó igual"
    corre(cuerpo())


def test_la_columna_de_comprobante_no_cuesta_una_foto(base):
    """«sí» o «no» sin traerse nada."""
    async def cuerpo():
        await base.transactions.insert_one(
            {"transaction_id": "con_lista", "proof_images": [FOTO]})
        await base.transactions.insert_one(
            {"transaction_id": "con_viejo", "proof_image": FOTO})
        await base.transactions.insert_one({"transaction_id": "sin_nada"})
        await base.transactions.insert_one(
            {"transaction_id": "lista_vacia", "proof_images": []})
        cuales = await las_fotos.cuales_tienen_foto(base, {})
        assert cuales == {"con_lista", "con_viejo"}
    corre(cuerpo())


def test_el_campo_viejo_y_el_nuevo_cuentan_los_dos(base):
    """Conviven: `proof_image` es el de antes y `proof_images` el de ahora.
    Mirar sólo uno es el defecto que se está arreglando."""
    assert "proof_image" in las_fotos.LAS_FOTOS
    assert "proof_images" in las_fotos.LAS_FOTOS


def test_sin_las_fotos_es_de_lo_prohibido_y_las_nombra_a_todas():
    assert set(las_fotos.SIN_LAS_FOTOS) == set(las_fotos.LAS_FOTOS)
    assert all(v == 0 for v in las_fotos.SIN_LAS_FOTOS.values())


def test_nadie_le_pasa_a_la_base_la_constante_compartida():
    """`SIN_LAS_FOTOS` se pide con `sin_las_fotos()`, que da una copia.

    El doble de Mongo de los tests le agrega `_id` a la proyección que
    recibe. Tres rutas le pasaban la constante misma, y el primer test que
    las llamó la dejó cambiada para el resto: falló
    `test_sin_las_fotos_es_de_lo_prohibido_y_las_nombra_a_todas`, que no
    tenía nada que ver, y sólo si corría después.
    """
    raiz = pathlib.Path(__file__).resolve().parent.parent
    la_nombran = []
    for archivo in raiz.rglob("*.py"):
        rel = archivo.relative_to(raiz).as_posix()
        if rel.startswith(("tests/", "services/las_fotos.py")) or "/node_modules/" in rel:
            continue
        for n, linea in enumerate(archivo.read_text(encoding="utf-8").splitlines(), 1):
            if "SIN_LAS_FOTOS" in linea:
                la_nombran.append(f"{rel}:{n}  {linea.strip()}")
    assert not la_nombran, (
        "Estas líneas usan la constante compartida en vez de una copia:\n\n"
        + "\n".join(f"    {x}" for x in la_nombran)
        + "\n\nUsá `las_fotos.sin_las_fotos()`."
    )


def test_solo_es_de_lo_permitido_y_no_deja_pasar_fotos():
    p = las_fotos.solo("transaction_id", "status")
    assert p == {"_id": 0, "transaction_id": 1, "status": 1}
    for campo in las_fotos.LAS_FOTOS:
        assert campo not in p


# ══════════════════════════════════════════════════════════════════════════
# 4. La exportación a Excel, llamada de verdad
# ══════════════════════════════════════════════════════════════════════════
#
# El test de arriba (`test_exportar_no_se_trae_ni_una_foto`) escribe su propia
# consulta: prueba que `las_fotos.solo` no trae fotos, no que la ruta lo use.
# Por eso la ruta de `routes/misc.py` pudo quedarse con `{"_id": 0}`. Éste
# llama a la ruta y mira qué le pidió a la base.

class _BaseQueAnota:
    def __init__(self, base, anotadas):
        self._base, self._anotadas = base, anotadas

    def __getattr__(self, nombre):
        coleccion = getattr(self._base, nombre)
        if nombre != "transactions":
            return coleccion
        anotadas = self._anotadas

        class _Col:
            def find(self, filtro, proyeccion=None, *a, **k):
                anotadas.append(proyeccion)
                return coleccion.find(filtro, proyeccion, *a, **k)

            def __getattr__(self, otro):
                return getattr(coleccion, otro)
        return _Col()


def _exportar(base):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from models.user import User
    from routes import dependencies as deps
    from routes.misc import router
    app = FastAPI()
    app.include_router(router, prefix="/api")
    jefa = User(user_id="u_jefa", name="Jefa", email="jefa@ejemplo.test", role="super_admin")
    app.dependency_overrides[deps.get_super_admin] = lambda: jefa
    app.dependency_overrides[deps.get_current_user] = lambda: jefa
    return TestClient(app).get("/api/transactions/export")


def test_la_ruta_de_exportar_no_le_pide_fotos_a_la_base(base):
    from services.money import to_decimal128
    corre(base.transactions.insert_many([
        {"transaction_id": f"tx{i}", "display_id": f"R{i:06d}", "type": "withdrawal",
         "status": "completed", "amount_ris": to_decimal128("10.00"),
         "proof_image": FOTO, "proof_images": [FOTO, FOTO]} for i in range(3)]))
    anotadas = []
    usar_base(_BaseQueAnota(base, anotadas))
    r = _exportar(base)
    assert r.status_code == 200, r.text
    (proyeccion,) = anotadas
    for campo in las_fotos.LAS_FOTOS:
        assert campo not in proyeccion, f"la exportación pide «{campo}»"
    assert set(proyeccion) - {"_id"}, "`{_id: 0}` sola trae la orden entera"
    assert len(r.content) < 30 * 1024, f"el Excel pesa {len(r.content)} bytes"


def test_la_exportacion_escribe_el_dinero_guardado_en_decimal128(base):
    """Antes contestaba 500: «Cannot convert Decimal128 to Excel»."""
    from io import BytesIO
    from openpyxl import load_workbook
    from services.money import to_decimal128
    corre(base.transactions.insert_many([
        {"transaction_id": "tx_decimal", "display_id": "R000001", "type": "recharge",
         "status": "completed", "amount_ris": to_decimal128("1234.56"),
         "amount_ves": to_decimal128("45067.89")},
        # `display_id` en null: antes daba `None[:15]`.
        {"transaction_id": "tx_sin_numero", "display_id": None, "type": "withdrawal",
         "status": "pending", "amount_ris": 7.5},
    ]))
    r = _exportar(base)
    assert r.status_code == 200, r.text
    filas = list(load_workbook(BytesIO(r.content)).active.iter_rows(values_only=True))
    por_id = {f[0]: f for f in filas[1:]}
    assert por_id["R000001"][3] == 1234.56 and por_id["R000001"][4] == 45067.89
    assert por_id["tx_sin_numero"][3] == 7.5
