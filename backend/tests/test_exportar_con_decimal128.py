"""
tests/test_exportar_con_decimal128.py — Cada exportación a Excel con la plata
guardada como la guarda la aplicación: en Decimal128.

El Excel no sabe escribir un Decimal128: contesta «Cannot convert Decimal128 to
Excel» y la exportación entera da 500. La de órdenes de routes/misc.py se
arregló en el #182; su gemela del panel (admin_routes.py) seguía rota el 26 de
septiembre de 2026, y ningún test lo veía: los de Reportes y del reporte
contable siembran los montos en `float`, que es justo lo que no alcanza (ver
CLAUDE.md, «Lo que mongomock no puede ver»).

Cada test mira que el monto llegue a la hoja como NÚMERO, no sólo que no haya
500: un monto convertido a texto también «anda» y no se puede sumar.
"""
import asyncio
import io
import os
import sys
import types
from datetime import datetime, timezone

import pytest

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))
os.environ.setdefault("MONGO_URL", "mongodb://localhost:27017")
os.environ.setdefault("DB_NAME", "ris_test")

pytest.importorskip("mongomock_motor")

if "pywebpush" not in sys.modules:
    _stub = types.ModuleType("pywebpush")
    _stub.WebPushException = type("WebPushException", (Exception,), {})
    _stub.webpush = lambda *a, **k: None
    sys.modules["pywebpush"] = _stub

from conftest import ensenarle_decimal128_a_mongomock, usar_base      # noqa: E402
ensenarle_decimal128_a_mongomock()

from models.user import User                                         # noqa: E402
from services.money import to_decimal128                             # noqa: E402

JEFA = User(user_id="u_jefa", name="Jefa", email="jefa@ejemplo.test", role="super_admin")
AHORA = datetime.now(timezone.utc)
HOY = AHORA.strftime("%Y-%m-%d")


def corre(coro):
    return asyncio.run(coro)


@pytest.fixture
def base(monkeypatch):
    import mongomock_motor
    import admin_routes
    b = mongomock_motor.AsyncMongoMockClient()["ris_exportar_d128"]
    usar_base(b)
    # admin_routes abre su propia conexión al importarse; sin esto, la
    # exportación del panel buscaría un Mongo de verdad.
    monkeypatch.setattr(admin_routes, "db", b)
    return b


def _una_orden(base):
    corre(base.transactions.insert_one({
        "transaction_id": "tx_1", "display_id": "RIS-0001", "user_id": "u_1",
        "user_email": "cliente@ejemplo.test", "type": "withdrawal", "status": "completed",
        "amount_input": to_decimal128("100.50"), "amount_output": to_decimal128("4500.25"),
        "amount_ris": to_decimal128("100.50"), "amount_ves": to_decimal128("4500.25"),
        "rate": to_decimal128("44.78"), "currency_input": "RIS", "currency_output": "VES",
        "created_at": AHORA, "completed_at": AHORA,
        "beneficiary_data": {"full_name": "Beneficiario"},
    }))


def _valores(contenido: bytes) -> list:
    from openpyxl import load_workbook
    libro = load_workbook(io.BytesIO(contenido))
    return [c.value for hoja in libro.worksheets for fila in hoja.iter_rows() for c in fila]


def _cuerpo(respuesta) -> bytes:
    """El archivo de una `Response` o de una `StreamingResponse`."""
    if hasattr(respuesta, "body"):
        return respuesta.body

    async def juntar():
        return b"".join([p async for p in respuesta.body_iterator])
    return corre(juntar())


def test_LA_EXPORTACION_DEL_PANEL_ESCRIBE_LA_PLATA_EN_DECIMAL128(base):
    """La que seguía rota: contestaba 500 con una sola orden así."""
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    import admin_routes
    _una_orden(base)
    app = FastAPI()
    app.include_router(admin_routes.admin_router)
    app.dependency_overrides[admin_routes.get_admin_user] = lambda: JEFA
    r = TestClient(app).get("/api/admin/transactions/export")
    assert r.status_code == 200, r.text
    valores = _valores(r.content)
    assert 100.5 in valores and 4500.25 in valores, "los montos no llegaron como número"


def test_EL_REPORTE_CONTABLE_ESCRIBE_LA_PLATA_EN_DECIMAL128(base):
    from routes import accounting
    _una_orden(base)
    reporte = corre(accounting.get_accounting_report(route="brl_ves", period="month", date=None, admin=JEFA))
    assert reporte["rows"], "sin filas, el test pasaría sin haber escrito ningún monto"
    r = corre(accounting.export_accounting_excel(route="brl_ves", period="month", date=None, admin=JEFA))
    assert 100.5 in _valores(_cuerpo(r))


def test_LOS_REPORTES_DEL_PANEL_ESCRIBEN_LA_PLATA_EN_DECIMAL128(base):
    from routes.admin import reportes
    _una_orden(base)
    pedido = dict(desde=HOY, hasta=HOY, flujos=None, buscar=None, operador=None, monto_min=None,
                  monto_max=None, tz_min=0, limite=100, saltear=0, admin=JEFA)
    assert corre(reportes.generar_reporte(formato="json", **pedido))["filas"], "sin filas no se prueba nada"
    r = corre(reportes.generar_reporte(formato="xlsx", **pedido))
    assert 100.5 in _valores(_cuerpo(r))


def test_EL_BALANCE_DEL_LIBRO_ESCRIBE_LA_PLATA_EN_DECIMAL128(base):
    """Con una línea escrita por la aplicación, no a mano: así queda como en
    producción."""
    from routes import ledger_admin
    from services import ledger
    corre(ledger.record_ris_entry(user_id="u_1", movement_type="recarga_pix", amount="100.50",
                                  direction="credit", balance_before="0", balance_after="100.50",
                                  reference_kind="manual", reference_id="r1"))
    assert type(corre(base.ledger.find_one({}))["amount"]).__name__ == "Decimal128"
    r = corre(ledger_admin.ver_balance(desde=HOY, hasta=HOY, libro=None, tz_min=0, formato="xlsx", admin=JEFA))
    assert 100.5 in _valores(_cuerpo(r))
